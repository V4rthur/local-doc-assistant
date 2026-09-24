"""Document loaders for PDF, DOCX, and TXT files."""
import hashlib
from pathlib import Path
from typing import Iterator

import pdfplumber
from docx import Document as DocxDocument

from src.ingestion.models import Block, Document
from src.ingestion.normalize import normalize


def _hash_file(path: Path) -> str:
    """SHA-256 of file bytes — used for incremental re-indexing."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def _looks_like_real_table(table: list[list[str | None]]) -> bool:
    """Reject pdfplumber false positives (prose in single-cell or 2-col grids).

    A real table has:
      - At least 2 rows AND at least 2 columns
      - Multiple cells per row actually filled (not just col 0)
      - Cells that aren't paragraph-length prose
    """
    if not table or len(table) < 2:
        return False

    # Normalize: replace None with ""
    clean = [[(c or "").strip() for c in row] for row in table]

    n_cols = max(len(row) for row in clean)
    if n_cols < 2:
        return False

    # Count rows where at least 2 cells are non-empty
    multi_filled_rows = sum(
        1 for row in clean if sum(1 for c in row if c) >= 2
    )
    if multi_filled_rows < 2:
        return False  # It's really a single-column stream

    # Reject if any cell is huge (>300 chars) — that's a paragraph, not a table cell
    all_cells = [c for row in clean for c in row if c]
    if all_cells and max(len(c) for c in all_cells) > 300:
        return False

    # Reject if average cell length is paragraph-like (>80 chars)
    if all_cells and sum(len(c) for c in all_cells) / len(all_cells) > 80:
        return False

    return True


def _table_to_markdown(table: list[list[str | None]]) -> str:
    """Convert a pdfplumber table (list of rows) to markdown."""
    if not table or not table[0]:
        return ""
    clean = [[(c or "").strip().replace("\n", " ") for c in row] for row in table]
    # Normalize row widths (pdfplumber sometimes returns ragged rows)
    max_cols = max(len(row) for row in clean)
    clean = [row + [""] * (max_cols - len(row)) for row in clean]

    header = clean[0]
    body = clean[1:]

    md_lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * len(header)) + " |",
    ]
    md_lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(md_lines)


# ---------------- PDF ----------------

def load_pdf(path: Path, normalize_script: bool = True) -> Document:
    """Extract text + tables from a PDF, one Block per page-content-unit.

    Uses a two-pass approach:
      1. Find candidate tables and validate them (reject false positives).
      2. Extract text — subtract only the bboxes of REAL tables.
      3. Emit valid tables as markdown blocks.
    """
    blocks: list[Block] = []
    total_pages = 0
    combined_text_for_lang: list[str] = []

    with pdfplumber.open(path) as pdf:
        total_pages = len(pdf.pages)
        for page_num, page in enumerate(pdf.pages, start=1):
            # 1. Find candidate tables and keep only the real ones
            candidate_tables = page.find_tables()
            real_tables = []
            for t in candidate_tables:
                extracted = t.extract()
                if _looks_like_real_table(extracted):
                    real_tables.append((t, extracted))

            real_bboxes = [t.bbox for t, _ in real_tables]

            # 2. Text extraction — only subtract REAL table bboxes
            def _outside_tables(obj):
                x0, y0, x1, y1 = obj["x0"], obj["top"], obj["x1"], obj["bottom"]
                for tx0, ty0, tx1, ty1 in real_bboxes:
                    if x0 >= tx0 and x1 <= tx1 and y0 >= ty0 and y1 <= ty1:
                        return False
                return True

            text = page.filter(_outside_tables).extract_text() or ""
            if text.strip():
                norm_text, _ = normalize(text, "latin" if normalize_script else "preserve")
                if norm_text:
                    blocks.append(Block(kind="text", content=norm_text, page=page_num))
                    combined_text_for_lang.append(norm_text)

            # 3. Emit only validated tables
            for _, extracted in real_tables:
                md = _table_to_markdown(extracted)
                if md:
                    norm_md, _ = normalize(md, "latin" if normalize_script else "preserve")
                    blocks.append(Block(kind="table", content=norm_md, page=page_num))

    from src.ingestion.normalize import detect_script
    lang = detect_script("\n".join(combined_text_for_lang)[:2000]) if combined_text_for_lang else "mixed"

    return Document(
        source_path=str(path),
        filename=path.name,
        file_hash=_hash_file(path),
        blocks=blocks,
        total_pages=total_pages,
        language_hint=lang,
    )


# ---------------- DOCX ----------------

def load_docx(path: Path, normalize_script: bool = True) -> Document:
    """Extract paragraphs and tables from a Word document.

    DOCX has no real 'pages' — we approximate by grouping every N paragraphs.
    """
    doc = DocxDocument(str(path))
    blocks: list[Block] = []
    combined_text_for_lang: list[str] = []

    APPROX_PARAS_PER_PAGE = 30  # rough heuristic for citations
    para_count = 0

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        para_count += 1
        page = (para_count - 1) // APPROX_PARAS_PER_PAGE + 1
        norm_text, _ = normalize(text, "latin" if normalize_script else "preserve")
        if norm_text:
            blocks.append(Block(kind="text", content=norm_text, page=page))
            combined_text_for_lang.append(norm_text)

    # Tables in DOCX
    for table_idx, table in enumerate(doc.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        md = _table_to_markdown(rows)
        if md:
            norm_md, _ = normalize(md, "latin" if normalize_script else "preserve")
            # Approximate: put table on a page after the last text block seen
            page = max((b.page for b in blocks), default=1)
            blocks.append(Block(kind="table", content=norm_md, page=page))

    from src.ingestion.normalize import detect_script
    lang = detect_script("\n".join(combined_text_for_lang)[:2000]) if combined_text_for_lang else "mixed"

    return Document(
        source_path=str(path),
        filename=path.name,
        file_hash=_hash_file(path),
        blocks=blocks,
        total_pages=max((b.page for b in blocks), default=1),
        language_hint=lang,
    )


# ---------------- TXT ----------------

def load_txt(path: Path, normalize_script: bool = True) -> Document:
    """Load a plain text file."""
    raw = path.read_text(encoding="utf-8", errors="replace")
    norm_text, lang = normalize(raw, "latin" if normalize_script else "preserve")
    blocks = [Block(kind="text", content=norm_text, page=1)] if norm_text else []
    return Document(
        source_path=str(path),
        filename=path.name,
        file_hash=_hash_file(path),
        blocks=blocks,
        total_pages=1,
        language_hint=lang,
    )


# ---------------- Router ----------------

LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".txt": load_txt,
}


def load_document(path: Path, normalize_script: bool = True) -> Document:
    """Dispatch to the right loader based on file extension."""
    ext = path.suffix.lower()
    if ext not in LOADERS:
        raise ValueError(f"Unsupported file type: {ext} (supported: {list(LOADERS)})")
    return LOADERS[ext](path, normalize_script=normalize_script)


def iter_raw_documents(raw_dir: Path) -> Iterator[Path]:
    """Yield every supported file under raw_dir."""
    for path in raw_dir.rglob("*"):
        if path.is_file() and path.suffix.lower() in LOADERS:
            yield path